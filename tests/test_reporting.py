from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook

from perfmon.config.schema import Targets
from perfmon.reporting import REPORTERS
from perfmon.reporting.console import format_summary
from perfmon.reporting.excel import build_workbook
from perfmon.reporting.summary import (
    FailureGroup,
    Iterations,
    Measure,
    Phase,
    ResponseTimes,
    RunSummary,
    StepKey,
    TimeSlice,
    target_misses,
)

LOGIN = StepKey("flow", "setup/login", "POST")
ORDER = StepKey("flow", "create_order", "POST")


def measure(times: list[float], failures: int = 0) -> Measure:
    result = Measure()
    for index, milliseconds in enumerate(times):
        result.record(milliseconds, failed=index < failures)
    return result


def phase(users: int, started: float, order_times: list[float], failures: int, verdict: str) -> Phase:
    order = measure(order_times, failures)
    total = measure(order_times, failures)
    iterations = Iterations(total=len(order_times), incomplete=failures, steps_not_sent=failures)
    return Phase(users, started, users / 2, 30.0, total, {ORDER: order}, iterations, verdict)


def summary(**changes) -> RunSummary:
    healthy = phase(10, 0.0, [200.0] * 60, 0, "healthy")
    strained = phase(20, 30.0, [200.0] * 90 + [1200.0] * 10, 6, "p95 was 1200 ms (threshold 1000 ms)")
    order = measure([200.0] * 150 + [1200.0] * 10, 6)
    iterations = Iterations(total=160, incomplete=6, steps_not_sent=6, setups_failed=1)
    for _ in range(154):
        iterations.durations.add(450.0)
    values = {
        "project": "shop",
        "mode": "limit",
        "host": "https://api.example.com",
        "started": datetime(2026, 10, 11, 20, 30, 5, tzinfo=UTC),
        "duration_seconds": 60.0,
        "settings": {"Users in the first step": 10, "Scenarios": "flow (weight 1)"},
        "findings": [
            "Limit reached at 20 users: p95 was 1200 ms (threshold 1000 ms).",
            "Every target was met up to 10 users.",
        ],
        "steps": {ORDER: order, LOGIN: measure([80.0] * 20)},
        "targets": {ORDER: Targets(p95_ms=500, error_percent=1), LOGIN: Targets()},
        "failures": [FailureGroup(ORDER, "unexpected status", 503, "expected status 201, got 503", 6, 41.5, 20)],
        "iterations": {"flow": iterations},
        "phases": [healthy, strained],
        "timeline": [TimeSlice(0, 30, 10, measure([200.0] * 60)), TimeSlice(30, 30, 20, measure([200.0] * 100, 6))],
    }
    return RunSummary(**(values | changes))


def rows(sheet) -> list[tuple]:
    return [tuple(cell.value for cell in row) for row in sheet.iter_rows()]


def test_response_times_keep_counts_not_samples():
    times = ResponseTimes()
    for milliseconds in [12.4, 12.6, 148.0, 152.0, 1234.0]:
        times.add(milliseconds)
    assert (times.count, times.minimum, times.maximum) == (5, 12.4, 1234.0)
    assert len(times._counts) == 4  # 12, 13, 150 (twice), 1200
    assert (times.percentile(0.5), times.percentile(0.8), times.percentile(1.0)) == (150, 150, 1200)

    other = ResponseTimes()
    other.add(5.0)
    times.merge(other)
    assert (times.count, times.minimum, times.percentile(0.1)) == (6, 5.0, 5)
    assert ResponseTimes().percentile(0.95) == 0


def test_targets_are_judged_per_measure():
    slow = measure([100.0] * 90 + [900.0] * 10, failures=3)
    assert target_misses(slow, Targets(p95_ms=1000, p99_ms=1000, error_percent=5)) == []
    assert target_misses(slow, Targets(p95_ms=500, p99_ms=800, error_percent=1)) == [
        "p95 900 ms > 500 ms",
        "p99 900 ms > 800 ms",
        "3.0% failed > 1%",
    ]
    assert target_misses(Measure(), Targets(p95_ms=1)) == []  # nothing measured, nothing missed
    assert Targets(p95_ms=300).over(Targets(p95_ms=800, error_percent=1)) == Targets(p95_ms=300, error_percent=1)


def test_summary_knows_which_steps_missed_their_targets():
    run = summary()
    assert run.requests == 180 and run.failed_requests == 6 and run.has_targets
    assert run.target_verdict(LOGIN, run.steps[LOGIN]) == ""  # no targets set
    assert run.target_verdict(ORDER, run.steps[ORDER]) == "missed: p95 1200 ms > 500 ms; 3.8% failed > 1%"
    assert list(run.steps_missing_targets()) == [ORDER]
    assert run.steps_missing_targets(run.phases[0].steps) == {}


def test_workbook_has_a_dashboard_and_a_sheet_per_table():
    workbook = build_workbook(summary())
    assert workbook.sheetnames == [
        "Dashboard", "Load steps", "Steps by load", "Steps", "Iterations", "Failures", "Timeline", "Settings", "Chart data",
    ]  # fmt: skip

    facts = dict(rows(workbook["Dashboard"])[:13])
    assert (facts["Project"], facts["Mode"], facts["Requests"], facts["Failed requests"]) == ("shop", "limit", 180, 6)
    assert facts["Started (UTC)"] == datetime(2026, 10, 11, 20, 30, 5) and facts["Incomplete iterations"] == 6
    assert facts["Findings"].startswith("Limit reached at 20 users")

    load_steps = rows(workbook["Load steps"])
    assert load_steps[0][:7] == (
        "Users", "Duration (s)", "Requests", "Requests/s", "Intended iterations/s", "Achieved iterations/s", "Failed requests (%)",
    )  # fmt: skip
    assert load_steps[1] == (10, 30, 60, 2, 5, 2, 0, 200, 200, 200, 200, 60, 0, 0, 0, "met", "healthy")
    assert load_steps[2][:11] == (20, 30, 100, 100 / 30, 10, 100 / 30, 6, 200, 1200, 1200, 1200)
    assert load_steps[2][15:] == ("missed by 1 step(s)", "p95 was 1200 ms (threshold 1000 ms)")

    by_load = rows(workbook["Steps by load"])
    assert by_load[2] == (
        20,
        "flow",
        "create_order",
        "POST",
        100,
        100 / 30,
        6,
        200,
        1200,
        1200,
        1200,
        "missed: p95 1200 ms > 500 ms; 6.0% failed > 1%",
    )

    steps = rows(workbook["Steps"])
    assert steps[1][:4] == ("flow", "create_order", "POST", 160) and steps[1][13:] == (
        500,
        None,
        1,
        "missed: p95 1200 ms > 500 ms; 3.8% failed > 1%",
    )
    assert steps[2][:4] == ("flow", "setup/login", "POST", 20) and steps[2][13:] == (None, None, None, None)
    assert rows(workbook["Iterations"])[1] == ("flow", 160, 6, 3.75, 6, 1, 450, 450, 450)
    assert rows(workbook["Failures"])[1] == (
        "flow",
        "create_order",
        "POST",
        "unexpected status",
        503,
        6,
        41.5,
        20,
        "expected status 201, got 503",
    )
    assert rows(workbook["Timeline"])[2] == (30, 20, 100, 100 / 30, 6, 200, 200)
    assert rows(workbook["Settings"]) == [
        ("Setting", "Value"),
        ("Users in the first step", 10),
        ("Scenarios", "flow (weight 1)"),
    ]


def test_tables_are_usable_as_they_open():
    sheet = build_workbook(summary())["Steps"]
    assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref == "A1:Q3"
    assert sheet["A1"].font.bold and sheet["H2"].number_format == "0" and sheet["G2"].number_format == "0.0"


def chart_titles(sheet) -> list[str]:
    return ["".join(run.t for paragraph in chart.title.tx.rich.p for run in paragraph.r) for chart in sheet._charts]


def test_a_stepped_run_charts_the_load_steps_and_the_run_over_time():
    assert chart_titles(build_workbook(summary())["Dashboard"]) == [
        "Response time by load (ms)",
        "Failed requests and incomplete iterations by load (%)",
        "Iterations per second by load, intended and achieved",
        "p95 by load, per step (ms)",
        "Response time over the run (ms)",
        "Virtual users over the run",
        "Requests per second over the run",
        "Failed requests over the run (%)",
    ]
    assert rows(build_workbook(summary())["Chart data"]) == [("Users", "POST flow/create_order"), (10, 200), (20, 1200)]


def test_a_baseline_charts_the_steps_against_their_targets():
    run = summary(mode="baseline", phases=[phase(10, 0.0, [200.0] * 60, 0, "")], failures=[])
    workbook = build_workbook(run)
    assert workbook.sheetnames == ["Dashboard", "Steps", "Iterations", "Failures", "Timeline", "Settings", "Chart data"]
    assert chart_titles(workbook["Dashboard"])[:2] == [
        "p95 per step, slowest first (ms)",
        "Response time over the run (ms)",
    ]
    assert rows(workbook["Chart data"]) == [
        ("Step", "p95 (ms)", "Target p95 (ms)"),
        ("POST flow/create_order", 1200, 500),
        ("POST flow/setup/login", 80, None),
    ]
    assert rows(workbook["Failures"]) == [
        (
            "Scenario",
            "Step",
            "Method",
            "Kind",
            "Status",
            "Occurrences",
            "First seen (s into run)",
            "First seen at users",
            "Message",
        )
    ]


def test_excel_reporter_names_the_file_after_project_mode_and_start(tmp_path: Path):
    reporter = REPORTERS["excel"](tmp_path / "out")
    message = reporter.on_end(summary())
    target = tmp_path / "out" / "shop-limit-20261011-203005.xlsx"
    assert message == f"Excel report written to {target}"
    assert load_workbook(target).sheetnames[0] == "Dashboard"


def test_console_text_has_findings_steps_failures_and_iterations():
    text = format_summary(summary())
    assert text.splitlines()[:2] == [
        "Limit reached at 20 users: p95 was 1200 ms (threshold 1000 ms).",
        "Every target was met up to 10 users.",
    ]
    lines = {line.split("  ")[0]: line.split() for line in text.splitlines()}
    assert lines["POST flow/create_order"][2:8] == ["160", "3.8", "200", "1200", "1200", "1200"]
    assert " ".join(lines["POST flow/create_order"][8:]) == "missed: p95 1200 ms > 500 ms; 3.8% failed > 1%"
    assert lines["POST flow/setup/login"][2:] == ["20", "0.0", "80", "80", "80", "80", "-"]
    assert (
        "6 x POST flow/create_order: expected status 201, got 503 [unexpected status 503; first after 42 s at 20 users]"
        in text
    )
    assert "flow: 1 virtual user(s) stopped because their setup failed" in text and "Iteration p50 ms" in text
