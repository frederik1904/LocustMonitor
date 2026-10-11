import re
import threading
import time
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

import pytest
from openpyxl import load_workbook

from helpers import write
from perfmon.cli import main
from perfmon.config.schema import BaselineSettings, LimitSettings, StopWhen
from perfmon.modes.baseline import Baseline
from perfmon.modes.limit import Limit, _breach
from perfmon.modes.load import run_load
from perfmon.project import open_project
from perfmon.reporting.console import format_summary
from perfmon.reporting.summary import Measure


class _Api(BaseHTTPRequestHandler):
    """Answers /ok at once; /busy takes 30 ms and fails while more than two requests are in flight."""

    in_flight = 0

    def do_GET(self):
        status = 200
        if self.path == "/busy":
            _Api.in_flight += 1
            status = 500 if _Api.in_flight > 2 else 200
            time.sleep(0.03)
            _Api.in_flight -= 1
        self.send_response(status)
        self.send_header("Content-Length", "2")
        self.end_headers()
        self.wfile.write(b"{}")

    def log_message(self, *args):
        pass


@pytest.fixture
def project(tmp_path: Path) -> Path:
    server = ThreadingHTTPServer(("127.0.0.1", 0), _Api)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    write(tmp_path, "project.yaml", f"host: http://127.0.0.1:{server.server_port}\n")
    write(tmp_path, "endpoints/ok/meta.yaml", "path: /ok\n")
    write(tmp_path, "endpoints/busy/meta.yaml", "path: /busy\n")
    write(tmp_path, "scenarios/light.yaml", "setup: [ok]\nsteps: [ok]\n")
    write(tmp_path, "scenarios/heavy.yaml", "weight: 5\nsteps: [ok]\n")
    write(tmp_path, "scenarios/crowded.yaml", "steps: [busy]\n")
    yield tmp_path
    server.shutdown()


def test_baseline_gives_every_scenario_its_users_and_paces_them(project: Path):
    loaded = open_project(project)
    mode = Baseline(BaselineSettings(users=2, duration=0.6, pacing=0.1))
    report = run_load(loaded, loaded.settings.host, mode, ["light", "heavy"], say=lambda line: None)
    assert report.failed_requests == 0
    counts = {
        line.split()[1]: int(line.split()[2]) for line in format_summary(report).splitlines() if line.startswith("GET ")
    }
    # 2 users x 0.6 s at one iteration per 0.1 s is about 12 each, whatever the scenario's weight.
    assert 8 <= counts["light/ok"] <= 14 and 8 <= counts["heavy/ok"] <= 14
    assert counts["light/setup/ok"] == 2  # once per virtual user


def test_limit_raises_the_load_until_the_error_threshold_is_crossed(project: Path):
    loaded = open_project(project)
    settings = LimitSettings(start_users=1, step_users=2, max_users=9, step_duration=0.5, pacing=0)
    lines: list[str] = []
    report = run_load(loaded, loaded.settings.host, Limit(settings), ["crowded"], say=lines.append)
    assert "1 users: " in lines[1] and " 0.0% failed" in lines[1]
    finding = format_summary(report).splitlines()[0]
    assert finding.startswith("Limit reached at ") and "of requests failed (threshold 5%)" in finding
    assert "the last healthy step was" in finding and "Failures:" in format_summary(report)


def test_limit_reports_when_nothing_gave_way(project: Path):
    loaded = open_project(project)
    settings = LimitSettings(start_users=1, step_users=1, max_users=2, step_duration=0.3, pacing=0.05)
    report = run_load(loaded, loaded.settings.host, Limit(settings), ["heavy"], say=lambda line: None)
    assert (
        format_summary(report).splitlines()[0]
        == "No limit found: still healthy at 2 users, the highest step within max_users."
    )


def measure(times: list[float], failures: int = 0) -> Measure:
    result = Measure()
    for index, milliseconds in enumerate(times):
        result.record(milliseconds, failed=index < failures)
    return result


def test_what_counts_as_over_the_limit():
    stop_when = StopWhen(error_percent=10, p95_ms=200)
    assert _breach(measure([100] * 9 + [150], failures=1), stop_when) is None
    assert _breach(measure([100] * 10, failures=2), stop_when) == "20.0% of requests failed (threshold 10%)"
    assert _breach(measure([100] * 19 + [900]), stop_when) is None  # one slow request in twenty stays under p95
    assert _breach(measure([100] * 18 + [900] * 2), stop_when) == "p95 was 900 ms (threshold 200 ms)"
    assert _breach(measure([]), stop_when) == "no request completed during the step"


def test_run_command_exit_codes(project: Path, capsys: pytest.CaptureFixture[str]):
    write(
        project,
        "project.yaml",
        (project / "project.yaml").read_text() + "modes: {baseline: {duration: 0.3, pacing: 0.1}}\n",
    )
    assert main(["run", "heavy", "--root", str(project)]) == 0
    out = capsys.readouterr().out
    assert "Baseline: 1 user(s) per scenario, 1 scenario(s), for 0.3 s" in out and "GET heavy/ok" in out

    write(project, "endpoints/ok/meta.yaml", "path: /ok\nvalidate: {status: 201}\n")
    assert main(["run", "heavy", "--root", str(project)]) == 1
    assert "heavy/ok: expected status 201, got 200" in capsys.readouterr().out


def test_limit_judges_on_status_alone_and_baseline_compares_bodies(project: Path):
    write(project, "endpoints/ok/meta.yaml", "path: /ok\nvalidate: {status: 200, body: exact}\n")
    write(project, "endpoints/ok/response.json", '{"never": "returned"}')
    loaded = open_project(project)
    quiet = {"say": lambda line: None}

    limit = Limit(LimitSettings(start_users=1, step_users=1, max_users=1, step_duration=0.3, pacing=0.05))
    assert run_load(loaded, loaded.settings.host, limit, ["heavy"], **quiet).failed_requests == 0

    baseline = Baseline(BaselineSettings(users=1, duration=0.3, pacing=0.05))
    report = run_load(loaded, loaded.settings.host, baseline, ["heavy"], **quiet)
    assert report.failed_requests == report.requests > 0 and "body: $.never is missing" in format_summary(report)


def test_incomplete_iterations_and_unsent_steps_are_counted_apart_from_request_failures(project: Path):
    write(project, "endpoints/gone/meta.yaml", "path: /gone\nvalidate: {status: 201}\n")
    write(project, "scenarios/broken.yaml", "steps: [ok, gone, ok, ok]\n")
    write(project, "scenarios/no_login.yaml", "setup: [gone]\nsteps: [ok]\n")
    loaded = open_project(project)
    lines: list[str] = []
    # Seven users over weights 5, 1 and 1 gives every scenario at least one.
    limit = Limit(LimitSettings(start_users=7, step_users=1, max_users=7, step_duration=0.4, pacing=0.1))
    report = run_load(loaded, loaded.settings.host, limit, ["broken", "heavy", "no_login"], say=lines.append)

    table = {line.split()[0]: line.split()[1:] for line in format_summary(report).splitlines() if line.strip()}
    iterations, incomplete, percent, not_sent = table["broken"][:4]
    # Every iteration of broken stops at its second step: two requests sent, two never sent.
    assert int(iterations) >= 2 and incomplete == iterations and percent == "100.0"
    assert int(not_sent) == 2 * int(iterations)
    assert table["heavy"][1:4] == ["0", "0.0", "0"]
    assert table["no_login"][:4] == ["0", "0", "0.0", "0"]
    assert "no_login: 1 virtual user(s) stopped because their setup failed" in format_summary(report)
    # The step's own line counts what finished within the step; iterations still running at its end come later.
    step = re.search(r"(\d+) of (\d+) iterations incomplete", lines[1])
    assert step and 1 <= int(step[1]) <= int(incomplete)


def test_run_command_writes_the_reports_it_is_asked_for(project: Path, capsys: pytest.CaptureFixture[str]):
    settings = "modes: {limit: {start_users: 1, step_users: 1, max_users: 2, step_duration: 0.3, pacing: 0.05}}\n"
    write(project, "project.yaml", (project / "project.yaml").read_text() + settings)
    out_dir = project / "out"
    arguments = ["run", "heavy", "--mode", "limit", "--root", str(project), "--report", "excel"]
    assert main([*arguments, "--report-dir", str(out_dir)]) == 0

    (report,) = out_dir.glob("*.xlsx")
    assert (
        report.name.startswith(f"{project.name}-limit-")
        and f"Excel report written to {report}" in capsys.readouterr().out
    )
    workbook = load_workbook(report)
    assert [row[0].value for row in workbook["Load steps"].iter_rows(min_row=2)] == [1, 2]
    assert workbook["Steps"]["A2"].value == "heavy" and workbook["Steps"]["B2"].value == "ok"


def test_targets_decide_whether_a_baseline_passes(project: Path, capsys: pytest.CaptureFixture[str]):
    host = (project / "project.yaml").read_text()
    quick = "modes: {baseline: {duration: 0.3, pacing: 0.05}}\n"

    write(project, "project.yaml", host + quick + "targets: {p95_ms: 5000, error_percent: 1}\n")
    write(project, "endpoints/ok/meta.yaml", "path: /ok\ntargets: {p95_ms: 0.001}\n")
    assert main(["run", "heavy", "--root", str(project)]) == 1
    out = capsys.readouterr().out
    assert "1 of 1 step(s) with targets missed them: GET heavy/ok" in out and "missed: p95 " in out

    # A step that fails every request passes as long as its own failure target allows it.
    write(project, "endpoints/ok/meta.yaml", "path: /ok\nvalidate: {status: 201}\ntargets: {error_percent: 100}\n")
    assert main(["run", "heavy", "--root", str(project)]) == 0
    assert "Every target was met (1 step(s) with targets)." in capsys.readouterr().out


def test_a_limit_run_says_up_to_which_load_the_targets_held(project: Path):
    write(project, "endpoints/busy/meta.yaml", "path: /busy\ntargets: {error_percent: 1}\n")
    loaded = open_project(project)
    settings = LimitSettings(start_users=1, step_users=3, max_users=7, step_duration=0.5, pacing=0)
    report = run_load(loaded, loaded.settings.host, Limit(settings), ["crowded"], say=lambda line: None)
    assert report.findings[1] in (
        "Every target was met up to 1 users.",
        "The targets were not met at any level of load.",
    )
    assert [step.users for step in report.phases][:2] == [1, 4] and report.phases[
        0
    ].intended_iterations_per_second is None
    assert report.timeline and report.failures[0].kind == "unexpected status" and report.failures[0].status == 500
