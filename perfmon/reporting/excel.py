"""The run as a workbook: a dashboard with charts first, then one sheet per table behind it."""

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet

from perfmon.reporting.base import RequestSample
from perfmon.reporting.summary import RunSummary, StepKey

_WHOLE = "0"
_ONE_DECIMAL = "0.0"
_TWO_DECIMALS = "0.00"

# Chart colours, used in this fixed order; checked for colour-blind separation as a set.
_SERIES = ("2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "008300")
_REFERENCE = "898781"  # a target drawn next to a measurement: neutral, so it does not read as another series
_GRID = "E1E0D9"
_CHART_WIDTH = 17.0  # cm
_CHART_HEIGHT = 9.0
_ROWS_PER_CHART = 20
_MOST_LINES = len(_SERIES)  # more lines than this in one chart cannot be told apart
_MOST_BARS = 15

Columns = list[tuple[str, str | None]]


class ExcelReporter:
    """Writes the run to one .xlsx file in a folder, named after the project, mode and start time."""

    def __init__(self, folder: Path) -> None:
        self._folder = folder

    def on_request(self, sample: RequestSample) -> None:
        pass

    def on_end(self, summary: RunSummary) -> str | None:
        self._folder.mkdir(parents=True, exist_ok=True)
        target = self._folder / f"{summary.project}-{summary.mode}-{summary.started:%Y%m%d-%H%M%S}.xlsx"
        build_workbook(summary).save(target)
        return f"Excel report written to {target}"


def build_workbook(summary: RunSummary) -> Workbook:
    workbook = Workbook()
    del workbook["Sheet"]  # the empty sheet a new workbook starts with
    dashboard = workbook.create_sheet("Dashboard")
    stepped = len(summary.phases) > 1 or summary.mode == "limit"
    load_steps = _load_steps_sheet(workbook.create_sheet("Load steps"), summary) if stepped else None
    if stepped:
        _steps_by_load_sheet(workbook.create_sheet("Steps by load"), summary)
    _steps_sheet(workbook.create_sheet("Steps"), summary)
    _iterations_sheet(workbook.create_sheet("Iterations"), summary)
    _failures_sheet(workbook.create_sheet("Failures"), summary)
    timeline = _timeline_sheet(workbook.create_sheet("Timeline"), summary)
    _settings_sheet(workbook.create_sheet("Settings"), summary)
    chart_data = workbook.create_sheet("Chart data")

    charts = []
    if load_steps is not None:
        charts += _load_charts(load_steps, len(summary.phases), paced=summary.phases[0].intended_iterations_per_second)
        charts += _slowest_by_load_chart(chart_data, summary)
    else:
        charts += _slowest_steps_chart(chart_data, summary)
    charts += _timeline_charts(timeline, len(summary.timeline), with_users=stepped)
    _dashboard_sheet(dashboard, summary, charts)
    return workbook


# --- dashboard ---


def _dashboard_sheet(sheet: Worksheet, summary: RunSummary, charts: list[Any]) -> None:
    iterations = summary.iterations.values()
    failed = 100 * summary.failed_requests / summary.requests if summary.requests else 0.0
    facts: list[tuple[str, Any, str | None]] = [
        ("Project", summary.project, None),
        ("Mode", summary.mode, None),
        ("Host", summary.host, None),
        ("Started (UTC)", summary.started.replace(tzinfo=None), "yyyy-mm-dd hh:mm:ss"),
        ("Duration (s)", summary.duration_seconds, _ONE_DECIMAL),
        ("Requests", summary.requests, _WHOLE),
        ("Failed requests", summary.failed_requests, _WHOLE),
        ("Failed requests (%)", failed, _ONE_DECIMAL),
        ("Iterations", sum(count.total for count in iterations), _WHOLE),
        ("Incomplete iterations", sum(count.incomplete for count in iterations), _WHOLE),
        ("Steps not sent", sum(count.steps_not_sent for count in iterations), _WHOLE),
    ]
    facts += [("Findings" if index == 0 else "", finding, None) for index, finding in enumerate(summary.findings)]
    for label, value, number_format in facts:
        sheet.append([label, value])
        sheet.cell(sheet.max_row, 1).font = Font(bold=True)
        sheet.cell(sheet.max_row, 2).alignment = Alignment(horizontal="left")
        if number_format:
            sheet.cell(sheet.max_row, 2).number_format = number_format
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 70
    # Printed or exported to PDF, the dashboard is one page wide.
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    # Two charts side by side under the facts: the left one over columns A-B, the right one from column C.
    top = sheet.max_row + 2
    for index, chart in enumerate(charts):
        sheet.add_chart(chart, f"{'A' if index % 2 == 0 else 'C'}{top + (index // 2) * _ROWS_PER_CHART}")


# --- tables ---


def _load_steps_sheet(sheet: Worksheet, summary: RunSummary) -> Worksheet:
    columns: Columns = [
        ("Users", _WHOLE),
        ("Duration (s)", _ONE_DECIMAL),
        ("Requests", _WHOLE),
        ("Requests/s", _ONE_DECIMAL),
        ("Intended iterations/s", _TWO_DECIMALS),
        ("Achieved iterations/s", _TWO_DECIMALS),
        ("Failed requests (%)", _ONE_DECIMAL),
        ("p50 (ms)", _WHOLE),
        ("p95 (ms)", _WHOLE),
        ("p99 (ms)", _WHOLE),
        ("Max (ms)", _WHOLE),
        ("Iterations", _WHOLE),
        ("Incomplete", _WHOLE),
        ("Incomplete iterations (%)", _ONE_DECIMAL),
        ("Steps not sent", _WHOLE),
        ("Targets", None),
        ("Result", None),
    ]
    rows = []
    for phase in summary.phases:
        missed = summary.steps_missing_targets(phase.steps)
        targets = None if not summary.has_targets else f"missed by {len(missed)} step(s)" if missed else "met"
        rows.append(
            (
                phase.users,
                phase.duration_seconds,
                phase.total.requests,
                phase.requests_per_second,
                phase.intended_iterations_per_second,
                phase.iterations_per_second,
                phase.total.failure_percent,
                phase.total.p50,
                phase.total.p95,
                phase.total.p99,
                phase.total.times.maximum,
                phase.iterations.total,
                phase.iterations.incomplete,
                phase.iterations.incomplete_percent,
                phase.iterations.steps_not_sent,
                targets,
                phase.verdict,
            )
        )
    _table(sheet, columns, rows)
    return sheet


def _steps_by_load_sheet(sheet: Worksheet, summary: RunSummary) -> None:
    columns: Columns = [
        ("Users", _WHOLE),
        ("Scenario", None),
        ("Step", None),
        ("Method", None),
        ("Requests", _WHOLE),
        ("Requests/s", _ONE_DECIMAL),
        ("Failed (%)", _ONE_DECIMAL),
        ("p50 (ms)", _WHOLE),
        ("p95 (ms)", _WHOLE),
        ("p99 (ms)", _WHOLE),
        ("Max (ms)", _WHOLE),
        ("Targets", None),
    ]
    rows = [
        (
            phase.users,
            key.scenario,
            key.step,
            key.method,
            measure.requests,
            measure.requests / phase.duration_seconds if phase.duration_seconds else 0.0,
            measure.failure_percent,
            measure.p50,
            measure.p95,
            measure.p99,
            measure.times.maximum,
            summary.target_verdict(key, measure) or None,
        )
        for phase in summary.phases
        for key, measure in sorted(phase.steps.items())
    ]
    _table(sheet, columns, rows)


def _steps_sheet(sheet: Worksheet, summary: RunSummary) -> None:
    columns: Columns = [
        ("Scenario", None),
        ("Step", None),
        ("Method", None),
        ("Requests", _WHOLE),
        ("Requests/s", _ONE_DECIMAL),
        ("Failures", _WHOLE),
        ("Failed (%)", _ONE_DECIMAL),
        ("p50 (ms)", _WHOLE),
        ("p95 (ms)", _WHOLE),
        ("p99 (ms)", _WHOLE),
        ("Average (ms)", _WHOLE),
        ("Min (ms)", _WHOLE),
        ("Max (ms)", _WHOLE),
        ("Target p95 (ms)", _WHOLE),
        ("Target p99 (ms)", _WHOLE),
        ("Target failed (%)", _ONE_DECIMAL),
        ("Targets", None),
    ]
    rows = []
    for key, measure in summary.steps.items():
        targets = summary.targets.get(key)
        rows.append(
            (
                key.scenario,
                key.step,
                key.method,
                measure.requests,
                measure.requests / summary.duration_seconds if summary.duration_seconds else 0.0,
                measure.failures,
                measure.failure_percent,
                measure.p50,
                measure.p95,
                measure.p99,
                measure.times.average,
                measure.times.minimum,
                measure.times.maximum,
                targets.p95_ms if targets else None,
                targets.p99_ms if targets else None,
                targets.error_percent if targets else None,
                summary.target_verdict(key, measure) or None,
            )
        )
    _table(sheet, columns, rows)


def _iterations_sheet(sheet: Worksheet, summary: RunSummary) -> None:
    columns: Columns = [
        ("Scenario", None),
        ("Iterations", _WHOLE),
        ("Incomplete", _WHOLE),
        ("Incomplete (%)", _ONE_DECIMAL),
        ("Steps not sent", _WHOLE),
        ("Users stopped in setup", _WHOLE),
        ("Complete iteration p50 (ms)", _WHOLE),
        ("Complete iteration p95 (ms)", _WHOLE),
        ("Complete iteration max (ms)", _WHOLE),
    ]
    rows = [
        (
            name,
            count.total,
            count.incomplete,
            count.incomplete_percent,
            count.steps_not_sent,
            count.setups_failed,
            count.durations.percentile(0.50),
            count.durations.percentile(0.95),
            count.durations.maximum,
        )
        for name, count in sorted(summary.iterations.items())
    ]
    _table(sheet, columns, rows)


def _failures_sheet(sheet: Worksheet, summary: RunSummary) -> None:
    columns: Columns = [
        ("Scenario", None),
        ("Step", None),
        ("Method", None),
        ("Kind", None),
        ("Status", _WHOLE),
        ("Occurrences", _WHOLE),
        ("First seen (s into run)", _ONE_DECIMAL),
        ("First seen at users", _WHOLE),
        ("Message", None),
    ]
    rows = [
        (
            failure.key.scenario,
            failure.key.step,
            failure.key.method,
            failure.kind,
            failure.status,
            failure.occurrences,
            failure.first_seen_seconds,
            failure.first_seen_users,
            failure.message,
        )
        for failure in summary.failures
    ]
    _table(sheet, columns, rows)


def _timeline_sheet(sheet: Worksheet, summary: RunSummary) -> Worksheet:
    columns: Columns = [
        ("Time (s)", _WHOLE),
        ("Users", _WHOLE),
        ("Requests", _WHOLE),
        ("Requests/s", _ONE_DECIMAL),
        ("Failed (%)", _ONE_DECIMAL),
        ("p50 (ms)", _WHOLE),
        ("p95 (ms)", _WHOLE),
    ]
    rows = [
        (
            piece.start_seconds,
            piece.users,
            piece.measure.requests,
            piece.requests_per_second,
            piece.measure.failure_percent,
            piece.measure.p50,
            piece.measure.p95,
        )
        for piece in summary.timeline
    ]
    _table(sheet, columns, rows)
    return sheet


def _settings_sheet(sheet: Worksheet, summary: RunSummary) -> None:
    _table(sheet, [("Setting", None), ("Value", None)], list(summary.settings.items()))
    sheet.column_dimensions["B"].alignment = Alignment(horizontal="left")


def _table(sheet: Worksheet, columns: Columns, rows: list[tuple[Any, ...]]) -> None:
    """A header row and data rows, with a frozen bold header, filters and fitted column widths."""
    sheet.append([title for title, _ in columns])
    for row in rows:
        sheet.append(list(row))
    for index, (title, number_format) in enumerate(columns, start=1):
        sheet.cell(1, index).font = Font(bold=True)
        longest = max([len(title), *(len(_shown(row[index - 1])) for row in rows)])
        sheet.column_dimensions[get_column_letter(index)].width = min(longest + 3, 80)
        if number_format:
            for row_number in range(2, len(rows) + 2):
                sheet.cell(row_number, index).number_format = number_format
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def _shown(value: Any) -> str:
    """Roughly what a cell displays, to size its column: numbers are shown rounded."""
    if value is None:
        return ""
    return f"{value:.2f}" if isinstance(value, float) else str(value)


# --- charts ---
# Every chart has one value axis. Measures on different scales get a chart each, side by side.


def _load_charts(sheet: Worksheet, steps: int, paced: float | None) -> list[Any]:
    """How the API responded as the load was raised, from the Load steps sheet."""
    charts = [
        _line(sheet, steps, "Response time by load", "ms", columns=[8, 9, 10], x_title="Virtual users"),
        _bars(sheet, steps, "Failed requests and incomplete iterations by load", "%", columns=[7, 14]),
    ]
    if paced is not None:
        title = "Iterations per second by load, intended and achieved"
        charts.append(_line(sheet, steps, title, None, columns=[5, 6], x_title="Virtual users"))
    return charts


def _slowest_by_load_chart(sheet: Worksheet, summary: RunSummary) -> list[Any]:
    """p95 of the slowest steps at each level of load: which step bends first."""
    worst = {
        key: max(phase.steps[key].p95 for phase in summary.phases if key in phase.steps)
        for phase in summary.phases
        for key in phase.steps
    }
    chosen = sorted(worst, key=lambda key: -worst[key])[:_MOST_LINES]
    if not chosen:
        return []
    sheet.append(["Users", *(key.label for key in chosen)])
    for phase in summary.phases:
        sheet.append([phase.users, *(phase.steps[key].p95 if key in phase.steps else None for key in chosen)])
    title = "p95 by load, per step"
    if len(worst) > len(chosen):
        title = f"p95 by load, the {len(chosen)} slowest of {len(worst)} steps"
    columns = list(range(2, len(chosen) + 2))
    return [_line(sheet, len(summary.phases), title, "ms", columns=columns, x_title="Virtual users")]


def _slowest_steps_chart(sheet: Worksheet, summary: RunSummary) -> list[Any]:
    """p95 per step for a run at one level of load, slowest first, beside its target where one is set."""
    ranked: list[tuple[StepKey, float]] = sorted(
        ((key, measure.p95) for key, measure in summary.steps.items()), key=lambda item: -item[1]
    )[:_MOST_BARS]
    if not ranked:
        return []
    with_targets = any(summary.targets[key].p95_ms is not None for key, _ in ranked if key in summary.targets)
    sheet.append(["Step", "p95 (ms)", *(["Target p95 (ms)"] if with_targets else [])])
    for key, p95 in ranked:
        target = summary.targets[key].p95_ms if key in summary.targets else None
        sheet.append([key.label, p95, *([target] if with_targets else [])])
    title = "p95 per step, slowest first"
    if len(summary.steps) > len(ranked):
        title = f"p95 of the {len(ranked)} slowest of {len(summary.steps)} steps"
    chart = _bars(sheet, len(ranked), title, "ms", columns=[2, 3] if with_targets else [2], horizontal=True)
    if with_targets:
        chart.series[1].graphicalProperties.solidFill = _REFERENCE
    chart.height = max(_CHART_HEIGHT, 1.2 + 0.75 * len(ranked))
    return [chart]


def _timeline_charts(sheet: Worksheet, points: int, with_users: bool) -> list[Any]:
    """The run over time, from the Timeline sheet. The charts share their time axis so they can be read together."""
    if points < 2:
        return []
    time = {"x_title": "Seconds into the run", "markers": False}
    charts = [_line(sheet, points, "Response time over the run", "ms", columns=[6, 7], **time)]
    if with_users:
        charts.append(_line(sheet, points, "Virtual users over the run", None, columns=[2], **time))
    charts.append(_line(sheet, points, "Requests per second over the run", None, columns=[4], **time))
    charts.append(_line(sheet, points, "Failed requests over the run", "%", columns=[5], **time))
    return charts


def _line(
    sheet: Worksheet, rows: int, title: str, unit: str | None, *, columns: list[int], x_title: str, markers: bool = True
) -> LineChart:
    chart = LineChart()
    _fill(chart, sheet, rows, title, unit, columns)
    chart.x_axis.title = x_title
    for index, series in enumerate(chart.series):
        colour = _SERIES[index % len(_SERIES)]
        series.smooth = False
        series.graphicalProperties.line.solidFill = colour
        series.graphicalProperties.line.width = 19050  # 1.5 pt
        series.marker.symbol = "circle" if markers else "none"
        if markers:
            series.marker.size = 6
            series.marker.graphicalProperties.solidFill = colour
            series.marker.graphicalProperties.line.solidFill = "FFFFFF"  # a ring that separates overlapping points
    return chart


def _bars(
    sheet: Worksheet, rows: int, title: str, unit: str | None, *, columns: list[int], horizontal: bool = False
) -> BarChart:
    chart = BarChart()
    chart.type = "bar" if horizontal else "col"
    chart.gapWidth = 60
    _fill(chart, sheet, rows, title, unit, columns)
    if horizontal:
        # Categories run top to bottom in table order; keep the value axis at the bottom all the same.
        chart.x_axis.scaling.orientation = "maxMin"
        chart.y_axis.crosses = "max"
    else:
        chart.x_axis.title = "Virtual users"
    for index, series in enumerate(chart.series):
        series.graphicalProperties.solidFill = _SERIES[index % len(_SERIES)]
        series.graphicalProperties.line.solidFill = "FFFFFF"  # a gap between neighbouring bars
    return chart


def _fill(chart: Any, sheet: Worksheet, rows: int, title: str, unit: str | None, columns: list[int]) -> None:
    """Shared setup: data from columns of a table whose first column is the x axis, and quiet axes."""
    for column in columns:
        chart.add_data(Reference(sheet, min_col=column, min_row=1, max_row=rows + 1), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=rows + 1))
    chart.title = f"{title} ({unit})" if unit else title
    chart.width, chart.height = _CHART_WIDTH, _CHART_HEIGHT
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.y_axis.scaling.min = 0
    chart.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=_GRID, w=6350))
    if len(columns) == 1:
        chart.legend = None  # the title already names the one series
    else:
        # Many entries would squeeze the plot from below, so they go beside it instead.
        chart.legend.position = "b" if len(columns) <= 3 else "r"
